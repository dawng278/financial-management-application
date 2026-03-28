import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_report():
    doc = docx.Document()

    # Title
    title = doc.add_heading('BÁO CÁO ĐỒ ÁN: HỆ THỐNG QUẢN LÝ TÀI CHÍNH CÁ NHÂN', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 1. Giới thiệu tổng quan
    doc.add_heading('1. Giới thiệu tổng quan', level=1)
    doc.add_paragraph('Personal Finance Manager (PFM) là một ứng dụng máy tính (Desktop Application) giúp người dùng quản lý các giao dịch tài chính cá nhân một cách hiệu quả, minh bạch và chuyên nghiệp. Ứng dụng cung cấp các công cụ để theo dõi thu nhập, chi tiêu, lập ngân sách và phân tích dữ liệu qua các biểu đồ trực quan.')

    # 2. Công nghệ sử dụng
    doc.add_heading('2. Kiến trúc và Công nghệ', level=1)
    p = doc.add_paragraph()
    p.add_run('Hệ thống được xây dựng trên nền tảng .NET Framework/Core với các công nghệ chính:').bold = True
    doc.add_paragraph('Ngôn ngữ lập trình: C# (WinForms).', style='List Bullet')
    doc.add_paragraph('Kiến trúc: Phân lớp (N-tier architecture) bao gồm UI, BLL (Business Logic), DAL (Data Access Layer), và Common.', style='List Bullet')
    doc.add_paragraph('Cơ sở dữ liệu: MongoDB (Lưu trữ dữ liệu linh hoạt) và SQLite/File-based DB tùy cấu hình.', style='List Bullet')
    doc.add_paragraph('Giao diện (UI): Sử dụng thư viện HopeUI, ReaLTaiizor để tạo hiệu ứng Modern/Glassmorphism cao cấp.', style='List Bullet')
    doc.add_paragraph('Biểu đồ: LiveCharts để trực quan hóa dữ liệu chi tiêu và thu nhập.', style='List Bullet')

    # 3. Các chức năng chính
    doc.add_heading('3. Các chức năng chính', level=1)
    
    doc.add_heading('3.1 Quản lý người dùng và Xác thực', level=2)
    doc.add_paragraph('Ứng dụng hỗ trợ đăng ký tài khoản và đăng nhập bảo mật. Thông tin người dùng được mã hóa và lưu trữ an toàn.')

    doc.add_heading('3.2 Quản lý Giao dịch (Transactions)', level=2)
    doc.add_paragraph('Ghi chép các khoản Thu (Income) và Chi (Expense).', style='List Bullet')
    doc.add_paragraph('Phân loại giao dịch theo hạng mục (Category) và tài khoản (Account).', style='List Bullet')
    doc.add_paragraph('Hỗ trợ tìm kiếm, lọc và xem lịch sử giao dịch chi tiết.', style='List Bullet')

    doc.add_heading('3.3 Quản lý Danh mục và Tài khoản', level=2)
    doc.add_paragraph('Người dùng có thể tạo các "Bucket" (Hạng mục chi tiêu) tùy chỉnh.', style='List Bullet')
    doc.add_paragraph('Quản lý nhiều tài khoản ngân hàng, ví điện tử hoặc tiền mặt.', style='List Bullet')

    doc.add_heading('3.4 Báo cáo và Thống kê (Reports)', level=2)
    doc.add_paragraph('Biểu đồ Spend Density (Mật độ chi tiêu) theo thời gian.', style='List Bullet')
    doc.add_paragraph('Báo cáo tỷ lệ chi tiêu theo hạng mục (Category distribution).', style='List Bullet')
    doc.add_paragraph('Theo dõi biến động dòng tiền (Income vs Expense).', style='List Bullet')

    doc.add_heading('3.5 Thiết lập và Đa ngôn ngữ', level=2)
    doc.add_paragraph('Ứng dụng hỗ trợ chuyển đổi linh hoạt giữa tiếng Việt và tiếng Anh, giúp người dùng dễ dàng tiếp cận và sử dụng.')

    # 4. Đặc điểm nổi bật
    doc.add_heading('4. Đặc điểm nổi bật', level=1)
    doc.add_paragraph('Giao diện hiện đại (Modern UI): Thiết kế dựa trên các xu hướng mới, sử dụng bảng màu hài hòa và hiệu ứng chuyển động mượt mà.', style='List Bullet')
    doc.add_paragraph('Dữ liệu thời gian thực: Đồng bộ hóa dữ liệu nhanh chóng giữa giao diện và cơ sở dữ liệu.', style='List Bullet')
    doc.add_paragraph('Tính tùy biến cao: Người dùng có thể tự định nghĩa các hạng mục chi tiêu theo nhu cầu thực tế.', style='List Bullet')

    # 5. Kết luận
    doc.add_heading('5. Kết luận', level=1)
    doc.add_paragraph('Đồ án Personal Finance Manager đã hoàn thiện các chức năng cơ bản và nâng cao của một ứng dụng quản lý tài chính. Ứng dụng không chỉ có tính thực tiễn cao mà còn đảm bảo được tính thẩm mỹ và trải nghiệm người dùng tối ưu. Trong tương lai, hệ thống có thể phát triển thêm các tính năng dự báo chi tiêu bằng AI hoặc đồng bộ hóa đám mây.')

    # Footer/Sign
    doc.add_paragraph('\n\n--- Báo cáo được tạo tự động bởi hệ thống ---')

    doc.save('Bao_Cao_Do_An_PersonalFinanceManager.docx')
    print("Báo cáo đã được tạo thành công: Bao_Cao_Do_An_PersonalFinanceManager.docx")

create_report()
